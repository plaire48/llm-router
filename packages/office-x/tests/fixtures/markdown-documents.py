"""Synthetic corpus only: headings, merged tables, notes, sparse cells and OCR."""
import sys
from pathlib import Path
from datetime import datetime
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
d = Document()
d.add_heading("누비 매출 보고서", level=1)
p = d.add_paragraph("핵심 ")
p.add_run("원문 보존").bold = True
p.add_run(" 검증입니다.")
d.add_paragraph("첫 번째 확인", style="List Bullet")
t = d.add_table(rows=3, cols=2)
for row, values in zip(t.rows, [["항목", "매출"], ["서울", 1200], ["부산", 800]]):
    for cell, value in zip(row.cells, values):
        cell.text = str(value)
d.sections[0].header.paragraphs[0].text = "머리글 보존"
d.save(root / "report.docx")
w = Workbook()
s = w.active
s.title = "매출"
for row in [["지역", "금액"], ["서울", 1200], ["부산", 800]]:
    s.append(row)
s["B4"] = "=SUM(B2:B3)"
s["A7"] = "병합된 제목"
s.merge_cells("A7:C7")
s["H100"] = "먼 셀 값"
q = w.create_sheet("날짜")
q["A1"] = datetime(2026, 10, 14)
q["A1"].number_format = "yyyy-mm-dd"
w.save(root / "book.xlsx")
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = "누비 발표"
s.placeholders[1].text = "첫 번째 본문"
s.notes_slide.notes_text_frame.text = "발표자 노트 보존"
t = s.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
for r, row in enumerate([["지역", "금액"], ["서울", "1200"]]):
    for c, text in enumerate(row):
        t.cell(r, c).text = text
p.save(root / "deck.pptx")
font_paths = [Path("C:/Windows/Fonts/malgun.ttf"), Path("/System/Library/Fonts/AppleSDGothicNeo.ttc")]
font = next((path for path in font_paths if path.exists()), None)
if font:
    image = Image.new("RGB", (1500, 900), "white")
    draw = ImageDraw.Draw(image)
    large = ImageFont.truetype(str(font), 72)
    small = ImageFont.truetype(str(font), 48)
    draw.text((90, 90), "누비 문서 읽기", fill="black", font=large)
    draw.text((90, 250), "MARKDOWN EXPORT CHECK", fill="black", font=small)
    draw.text((90, 390), "매출 1200원", fill="black", font=large)
    image.save(root / "scanned.pdf", "PDF", resolution=150)
(root / "content.html").write_text('<h1>누비 HTML</h1><p><strong>원문 보존</strong></p><table><tr><td>서울</td><td>1200</td></tr></table>', encoding="utf-8")
